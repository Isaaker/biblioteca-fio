"""
Lectura del Excel maestro ("Índice General") de la Biblioteca FIO.

Adaptado de library/excel_import.py del proyecto Flask original: mismo
mapeo de columnas, misma hoja y misma fila de cabecera, para que el Excel
que ya usan los voluntarios funcione sin cambios. La diferencia es que
aquí no se construye un ORM `Book`: se genera un diccionario plano con
TODOS los campos (públicos e internos), y es config.py + safety_check.py
quien decide qué sale a la salida pública.

Columnas de la hoja "Indice General" (fila de cabecera = fila 2):
    A Clave                  -> reference
    B Título                 -> title
    C Subtítulo              -> subtitle
    D Autor                  -> author
    E NºEdición >1ª          -> edition
    F Lugar                  -> place
    G Editora                -> publisher
    H Año                    -> publication_year / publication_year_text
    I Tema                   -> topic
    J SubTema                -> subtopic
    K Especialidad           -> speciality
    L Especificaciones       -> specifications
    M Vigente/Histórico      -> status              (INTERNO)
    N Idioma                 -> language
    O Observaciones          -> notes               (INTERNO)
    P Páginas                -> num_pages
    Q ilustraciones          -> illustrations
    R cm                     -> size_cm
    S serie                  -> series
    T ISBN                   -> isbn
    U Repe 1                 -> repeated_copy_of    (INTERNO)
    V Donado por              -> donated_by          (INTERNO)
    W Medidas                -> measurements         (INTERNO)
    X Signatura Final        -> shelf_location       (INTERNO)
    Y Digitalizado (Sí/No)   -> is_digitized          [opcional, ver abajo]
    Z Enlace copia digital   -> digital_link          [opcional, ver abajo]

Las columnas Y/Z (digitalización) no existían en el Excel original de la
app Flask (allí ese dato se gestionaba a mano desde la web). Si tu Excel
maestro no las tiene, no pasa nada: is_digitized se queda en False y
digital_link en None para todos los libros; añádelas cuando quieras
empezar a publicar enlaces a copias digitales.
"""
from __future__ import annotations

from typing import Iterator

import openpyxl

from text_utils import normalize_language_field, normalize_topic_field

SHEET_NAME = "Indice General"
FIRST_DATA_ROW = 3


class ExcelReadError(Exception):
    """Error controlado al procesar el Excel de inventario."""


def _clean_str(value, max_len: int | None = None) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    return text[:max_len] if max_len else text


def _clean_int(value) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _clean_bool(value) -> bool:
    if value is None:
        return False
    text = str(value).strip().lower()
    return text in ("si", "sí", "yes", "true", "1", "x")


def read_inventory(excel_path: str) -> Iterator[dict]:
    """Genera un diccionario por libro con TODOS los campos del Excel
    (públicos e internos). No decide qué se publica: eso es
    responsabilidad de quien consume este generador (ver config.py).
    """
    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
    except Exception as exc:
        raise ExcelReadError(f"No se pudo leer el fichero Excel '{excel_path}': {exc}") from exc

    if SHEET_NAME not in wb.sheetnames:
        raise ExcelReadError(
            f"El Excel no contiene una hoja llamada '{SHEET_NAME}'. "
            f"Hojas encontradas: {', '.join(wb.sheetnames)}"
        )

    ws = wb[SHEET_NAME]

    for row in ws.iter_rows(min_row=FIRST_DATA_ROW):
        values = [cell.value for cell in row]
        while len(values) < 26:
            values.append(None)

        reference = _clean_str(values[0], 20)
        if not reference:
            continue  # fila vacía o sin clave ("Clave"): se ignora

        title = _clean_str(values[1], 255) or "Sin título"
        subtitle = _clean_str(values[2], 255)
        author = _clean_str(values[3], 255) or "Desconocido"
        edition = _clean_str(values[4], 50)
        place = _clean_str(values[5], 120)
        publisher = _clean_str(values[6], 255)

        year_raw = values[7]
        publication_year = _clean_int(year_raw)
        publication_year_text = _clean_str(year_raw, 20)

        topic = normalize_topic_field(_clean_str(values[8], 120))
        subtopic = normalize_topic_field(_clean_str(values[9], 120))
        speciality = _clean_str(values[10], 120)
        specifications = _clean_str(values[11], 255)
        status = _clean_str(values[12], 30)
        language = normalize_language_field(_clean_str(values[13], 60))
        notes = _clean_str(values[14], 2000)
        num_pages = _clean_int(values[15])
        illustrations = _clean_str(values[16], 255)
        size_cm = _clean_str(values[17], 20)
        series = _clean_str(values[18], 255)
        isbn = _clean_str(values[19], 20)
        repeated_copy_of = _clean_str(values[20], 20)
        donated_by = _clean_str(values[21], 255)
        measurements = _clean_str(values[22], 20)
        shelf_location = _clean_str(values[23], 120)
        is_digitized = _clean_bool(values[24])
        digital_link = _clean_str(values[25], 500)

        yield dict(
            reference=reference,
            title=title,
            subtitle=subtitle,
            author=author,
            edition=edition,
            place=place,
            publisher=publisher,
            publication_year=publication_year,
            publication_year_text=publication_year_text,
            topic=topic,
            subtopic=subtopic,
            speciality=speciality,
            specifications=specifications,
            language=language,
            num_pages=num_pages,
            illustrations=illustrations,
            size_cm=size_cm,
            series=series,
            isbn=isbn,
            is_digitized=is_digitized,
            digital_link=(digital_link if is_digitized else None),
            # --- internos, nunca deben llegar a una salida pública ---
            status=status,
            notes=notes,
            repeated_copy_of=repeated_copy_of,
            donated_by=donated_by,
            measurements=measurements,
            shelf_location=shelf_location,
        )
